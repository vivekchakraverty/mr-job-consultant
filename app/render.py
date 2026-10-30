"""ATS-friendly resume and cover letter files: TXT, DOCX, PDF.

ATS rules followed: one column, no tables/text boxes/images, contact details in the body (not the
page header), standard section headings, plain bullets, real selectable text, common fonts.
"""
import os
import re
from datetime import date
from pathlib import Path

HEADINGS = {"summary": "Summary", "skills": "Skills", "experience": "Experience", "projects": "Projects",
            "education": "Education", "certifications": "Certifications", "languages": "Languages"}


def _dates(entry: dict) -> str:
    return " – ".join(x for x in (entry.get("start"), entry.get("end")) if x)


def resume_blocks(r: dict) -> list[tuple[str, object]]:
    """The resume as an ordered list of (kind, payload) that every format renders the same way."""
    blocks: list[tuple[str, object]] = [("name", r.get("name") or "")]
    if r.get("headline"):
        blocks.append(("headline", r["headline"]))
    if r.get("contact"):
        blocks.append(("contact", " | ".join(r["contact"])))
    if r.get("summary"):
        blocks += [("heading", HEADINGS["summary"]), ("para", r["summary"])]
    if r.get("skills"):
        blocks += [("heading", HEADINGS["skills"]), ("para", ", ".join(r["skills"]))]
    if r.get("experience"):
        blocks.append(("heading", HEADINGS["experience"]))
        for e in r["experience"]:
            title = " – ".join(x for x in (e.get("title"), e.get("company")) if x)
            meta = " | ".join(x for x in (e.get("location"), _dates(e)) if x)
            blocks.append(("role", (title, meta)))
            blocks += [("bullet", b) for b in e.get("bullets", [])]
    if r.get("projects"):
        blocks.append(("heading", HEADINGS["projects"]))
        for p in r["projects"]:
            blocks.append(("role", (p.get("name", ""), "")))
            blocks += [("bullet", b) for b in p.get("bullets", [])]
    if r.get("education"):
        blocks.append(("heading", HEADINGS["education"]))
        for e in r["education"]:
            title = " – ".join(x for x in (e.get("degree"), e.get("institution")) if x)
            meta = " | ".join(x for x in (e.get("location"), _dates(e)) if x)
            blocks.append(("role", (title, meta)))
            blocks += [("bullet", d) for d in e.get("details", [])]
    for key in ("certifications", "languages"):
        if r.get(key):
            blocks.append(("heading", HEADINGS[key]))
            blocks += [("bullet", x) for x in r[key]]
    for o in r.get("other") or []:
        blocks.append(("heading", o["heading"]))
        blocks += [("bullet", x) for x in o["items"]]
    return blocks


def letter_blocks(name: str, contact: list[str], company: str, title: str, paragraphs: list[str]) -> list[tuple[str, object]]:
    blocks = [("name", name), ("contact", " | ".join(contact)), ("gap", None),
              ("para", date.today().strftime("%d %B %Y")), ("gap", None)]
    if company:
        blocks += [("para", f"{company} – Hiring team"), ("gap", None)]
    if title:
        blocks += [("para", f"Re: {title}"), ("gap", None)]
    blocks += [("para", "Dear Hiring Manager,"), ("gap", None)]
    for p in paragraphs:
        blocks += [("para", p), ("gap", None)]
    blocks += [("para", "Kind regards,"), ("para", name)]
    return blocks


# ---------- TXT ----------

def to_txt(blocks, path: Path) -> None:
    lines = []
    for kind, v in blocks:
        if kind == "name":
            lines.append(str(v).upper())
        elif kind in ("headline", "contact", "para"):
            lines.append(str(v))
        elif kind == "heading":
            lines += ["", str(v).upper()]
        elif kind == "role":
            title, meta = v
            lines += ["", title] + ([meta] if meta else [])
        elif kind == "bullet":
            lines.append(f"- {v}")
        elif kind == "gap":
            lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


# ---------- DOCX ----------

def to_docx(blocks, path: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches

    d = docx.Document()
    for section in d.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", bold=False, size=None, italic=False, align=None, space_before=0):
        p = d.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if align:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        return p

    for kind, v in blocks:
        if kind == "name":
            para(str(v), bold=True, size=18)
        elif kind == "headline":
            para(str(v), size=12)
        elif kind == "contact":
            para(str(v), size=10)
        elif kind == "heading":
            h = d.add_heading(str(v).upper(), level=2)  # real heading style: ATS parsers use it
            h.paragraph_format.space_before = Pt(10)
        elif kind == "role":
            title, meta = v
            para(title, bold=True, space_before=6)
            if meta:
                para(meta, italic=True, size=10)
        elif kind == "bullet":
            d.add_paragraph(str(v), style="List Bullet")
        elif kind == "para":
            para(str(v), align=WD_ALIGN_PARAGRAPH.LEFT)
        elif kind == "gap":
            para("")
    d.save(path)


# ---------- PDF ----------

FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
    ("C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf", "C:/Windows/Fonts/calibrii.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf"),
]


def to_pdf(blocks, path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=16)
    family, unicode_ok, bullet = "Helvetica", False, "-"
    for regular, bold, italic in FONT_CANDIDATES:
        if all(os.path.exists(p) for p in (regular, bold, italic)):
            pdf.add_font("Body", "", regular)
            pdf.add_font("Body", "B", bold)
            pdf.add_font("Body", "I", italic)
            family, unicode_ok, bullet = "Body", True, "•"
            break
    pdf.add_page()

    def clean(text: str) -> str:
        text = str(text)
        if unicode_ok:
            return text
        text = text.replace("–", "-").replace("—", "-").replace("•", "-").replace("’", "'").replace("“", '"').replace("”", '"')
        return text.encode("latin-1", "replace").decode("latin-1")

    width = pdf.w - pdf.l_margin - pdf.r_margin

    def write(text, style="", size=10.5, h=5.2, indent=0.0):
        pdf.set_font(family, style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, h, clean(text), new_x="LMARGIN", new_y="NEXT")

    for kind, v in blocks:
        if kind == "name":
            write(v, "B", 18, 9)
        elif kind == "headline":
            write(v, "", 12, 6)
        elif kind == "contact":
            write(v, "", 9.5, 5)
        elif kind == "heading":
            pdf.ln(3)
            write(str(v).upper(), "B", 11, 6)
            y = pdf.get_y()
            pdf.set_draw_color(160, 160, 160)
            pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
            pdf.ln(1.5)
        elif kind == "role":
            title, meta = v
            pdf.ln(1.5)
            write(title, "B", 10.5)
            if meta:
                write(meta, "I", 9.5, 4.8)
        elif kind == "bullet":
            pdf.set_font(family, "", 10.5)
            pdf.set_x(pdf.l_margin + 2)
            pdf.cell(4, 5.2, clean(bullet))
            pdf.multi_cell(width - 6, 5.2, clean(v), new_x="LMARGIN", new_y="NEXT")
        elif kind == "para":
            write(v)
        elif kind == "gap":
            pdf.ln(3)
    pdf.output(str(path))


WRITERS = {"pdf": to_pdf, "docx": to_docx, "txt": to_txt}


def safe_name(text: str) -> str:
    return re.sub(r"[^\w .-]+", "", text).strip()[:80] or "document"


def write_files(blocks, base: Path, formats: list[str]) -> list[Path]:
    out = []
    for fmt in formats:
        path = base.with_suffix(f".{fmt}")
        WRITERS[fmt](blocks, path)
        out.append(path)
    return out
