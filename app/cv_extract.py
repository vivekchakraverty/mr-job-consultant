"""Pull the contact details application forms always ask for out of an uploaded CV."""
import re
from pathlib import Path

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"(?<![\w/])(\+?\d[\d\s().-]{7,}\d)(?![\w/])")
LINKEDIN = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w%-]+/?", re.I)
WEBSITE = re.compile(r"(?:https?://)?(?:www\.)?(?:github\.com/[\w-]+|[\w-]+\.(?:dev|io|me|com|net|org)(?:/[\w-]*)?)", re.I)
NOT_A_NAME = re.compile(r"curriculum|vitae|resume|résumé|\bcv\b|profile|summary|contact|address|@|\d", re.I)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages[:6])
    if suffix == ".docx":
        import docx
        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    return ""  # .doc / .rtf are stored but not parsed


def _guess_name(lines: list[str]) -> str | None:
    for line in lines[:8]:
        words = line.split()
        if 2 <= len(words) <= 4 and not NOT_A_NAME.search(line) and all(w[:1].isupper() for w in words):
            return line.strip()
    return None


def _clean_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw)
    if not 8 <= len(digits) <= 15 or re.fullmatch(r"(19|20)\d{2}\D*(19|20)\d{2}", raw.strip()):
        return None  # too short/long, or a date range like "2019 - 2023"
    return " ".join(raw.split())


def extract_contact(text: str) -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    email = EMAIL.search(text)
    linkedin = LINKEDIN.search(text)
    phone = next((p for m in PHONE.finditer(text) if (p := _clean_phone(m.group(1)))), None)
    website = next((m.group(0) for m in WEBSITE.finditer(text)
                    if "linkedin" not in m.group(0).lower() and (not email or m.group(0) not in email.group(0))), None)
    return {
        "name": _guess_name(lines),
        "email": email.group(0) if email else None,
        "phone": phone,
        "linkedin": linkedin.group(0) if linkedin else None,
        "website": website,
    }
